import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Colors
    # Primary: #1E3A8A (Navy Blue), Secondary: #0284C7 (Light Blue), Text: #1F2937, Accent: #F3F4F6

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Real-Time Team Chat Application\n")
    run_title.bold = True
    run_title.font.size = Pt(24)
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    run_title.font.name = "Calibri"

    run_sub = p_title.add_run("Complete Architecture, Design Blueprint & Placement Guide\n")
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    run_sub.font.name = "Calibri"

    doc.add_paragraph("―" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1: Executive Summary & Placement Value
    h1 = doc.add_heading("1. Project Vision & Career Outcomes", level=1)
    h1.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p1 = doc.add_paragraph(
        "This project is a production-grade, distributed real-time chat application tailored for high-concurrency "
        "campus placement and technical interview discussions. Unlike basic single-file chat tutorials, this system "
        "incorporates enterprise architectural patterns including decoupled services, state synchronization across distributed instances, "
        "relational data modeling with raw SQL, and in-memory ephemeral caching."
    )
    p1.style.font.size = Pt(11)

    doc.add_heading("Key Learning Outcomes for Technical Interviews:", level=2)
    bullets = [
        ("Relational Integrity vs. Ephemeral Caching: ", "Deep mastery over when to use PostgreSQL (persistent chat history, user accounts, room hierarchies) versus Redis (presence tracking, typing status, socket session mappings)."),
        ("Horizontal Scaling with WebSockets: ", "Explaining how the Socket.io Redis Adapter utilizes Pub/Sub channels to broadcast messages across multiple backend Node.js processes seamlessly."),
        ("Stateless Auth in Stateful Protocols: ", "Implementing JWT verification during the initial HTTP handshake of a WebSocket connection to secure duplex channels."),
        ("Connection Lifecycle Resiliency: ", "Handling disconnections, client reconnects, multi-tab presence states, and heartbeat health checks.")
    ]
    for title, desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        r1 = bp.add_run(title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        bp.add_run(desc)

    # Section 2: Technology Stack
    h2 = doc.add_heading("2. Technology Stack & Component Responsibilities", level=1)
    h2.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    table_stack = doc.add_table(rows=1, cols=3)
    table_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_stack.autofit = False

    headers = ["Layer / Technology", "Specific Choice", "Architectural Responsibility"]
    col_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]

    hdr_cells = table_stack.rows[0].cells
    for i, name in enumerate(headers):
        hdr_cells[i].text = name
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], 140, 140, 160, 160)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10.5)

    stack_data = [
        ("Frontend Client", "React 19 + Vite (JS)", "Component lifecycle, client routing (react-router-dom), custom useSocket hook, optimistic UI updates."),
        ("Backend Server", "Node.js (LTS) + Express 5", "REST API endpoints, JWT authentication middleware, Zod payload validation, request orchestration."),
        ("Real-Time Engine", "Socket.io 4", "Bidirectional event messaging, room channels, connection lifecycle management."),
        ("Scale Adapter", "@socket.io/redis-adapter", "Redis Pub/Sub message broker to synchronize WebSocket events across multiple server instances."),
        ("Persistent DB", "PostgreSQL 16 (pg driver)", "Relational storage for users, rooms, memberships, and indexed message history. No ORM overhead."),
        ("In-Memory Store", "Redis 7", "Ephemeral state caching: active online sets, multi-socket tracking, debounced typing indicators."),
        ("DevOps / Setup", "Docker Compose", "Reproducible local containerization for PostgreSQL and Redis services.")
    ]

    for row_idx, data in enumerate(stack_data):
        row = table_stack.add_row()
        bg = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(data):
            cell = row.cells[i]
            cell.text = text
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9.5)
            if i == 0:
                p.runs[0].font.bold = True

    for row in table_stack.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    doc.add_paragraph()

    # Section 3: Feature Breakdown
    h3 = doc.add_heading("3. Core Feature Specifications", level=1)
    h3.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    features = [
        ("Authentication & Security", "User registration and login using salted bcrypt password hashing. Stateless JWTs verified on REST requests and during WebSocket handshakes. Input sanitization with Zod."),
        ("Chat Rooms & Channels", "Public and private channel creation, discovery, and membership management with many-to-many relational tables in PostgreSQL."),
        ("Real-Time Messaging", "Instant messaging with delivery acknowledgements, optimistic UI state updates, and automatic persistence into indexed PostgreSQL tables."),
        ("Presence Management", "Live online/offline user status tracking backed by Redis Sets. Handles multi-tab browser sessions so a user is only marked offline when all sockets disconnect."),
        ("Typing Indicators", "Transient typing event emissions with automatic debounce and server-side timeouts to prevent lingering typing badges."),
        ("Message History & Pagination", "Fast paginated history loading using composite SQL indexes (room_id, created_at DESC) for smooth reverse infinite scrolling.")
    ]

    for title, desc in features:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(f"{title}: ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        p.add_run(desc)

    # Section 4: Database Schema Design
    h4 = doc.add_heading("4. Database Schema (PostgreSQL 16 Pure SQL)", level=1)
    h4.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    sql_code = """-- 1. Users Table
CREATE TABLE users (
    id SERIAL PRIMARY KEY,
    username VARCHAR(50) UNIQUE NOT NULL,
    email VARCHAR(100) UNIQUE NOT NULL,
    password_hash VARCHAR(255) NOT NULL,
    avatar_url VARCHAR(255),
    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP
);

-- 2. Chat Rooms Table
CREATE TABLE rooms (
    id SERIAL PRIMARY KEY,
    name VARCHAR(100) UNIQUE NOT NULL,
    description TEXT,
    is_private BOOLEAN DEFAULT FALSE,
    created_by INTEGER REFERENCES users(id) ON DELETE SET NULL,
    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP
);

-- 3. Room Membership (Many-to-Many Join Table)
CREATE TABLE room_members (
    room_id INTEGER REFERENCES rooms(id) ON DELETE CASCADE,
    user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
    joined_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (room_id, user_id)
);

-- 4. Messages Table
CREATE TABLE messages (
    id SERIAL PRIMARY KEY,
    room_id INTEGER REFERENCES rooms(id) ON DELETE CASCADE,
    sender_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
    content TEXT NOT NULL,
    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP
);

-- Performance Indexes
CREATE INDEX idx_messages_room_created_at ON messages(room_id, created_at DESC);
CREATE INDEX idx_room_members_user_id ON room_members(user_id);"""

    p_code = doc.add_paragraph()
    set_cell_background(table_stack.rows[0].cells[0], "1E3A8A")
    r_code = p_code.add_run(sql_code)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Section 5: Real-Time Event Contracts
    h5 = doc.add_heading("5. Socket.io Real-Time Event Specifications", level=1)
    h5.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    table_events = doc.add_table(rows=1, cols=4)
    table_events.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_events.autofit = False

    ev_headers = ["Event Name", "Direction", "Payload", "Description"]
    ev_col_widths = [Inches(1.5), Inches(1.3), Inches(1.8), Inches(1.9)]

    hdr_cells2 = table_events.rows[0].cells
    for i, name in enumerate(ev_headers):
        hdr_cells2[i].text = name
        set_cell_background(hdr_cells2[i], "1E3A8A")
        set_cell_margins(hdr_cells2[i], 140, 140, 160, 160)
        p = hdr_cells2[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10)

    events_data = [
        ("room:join", "Client -> Server", "{ roomId }", "Joins client socket to the room channel"),
        ("room:leave", "Client -> Server", "{ roomId }", "Leaves the specified room channel"),
        ("message:send", "Client -> Server", "{ roomId, content }", "Persists message in SQL and triggers broadcast"),
        ("message:received", "Server -> Client", "{ id, roomId, sender, content, time }", "Dispatched to room members"),
        ("typing:start", "Client -> Server", "{ roomId }", "Registers user typing state in Redis"),
        ("typing:stop", "Client -> Server", "{ roomId }", "Clears user typing state in Redis"),
        ("typing:update", "Server -> Client", "{ roomId, typingUsers }", "Broadcasts currently typing users list"),
        ("presence:update", "Server -> Client", "{ userId, status }", "Broadcasts user online/offline status")
    ]

    for row_idx, data in enumerate(events_data):
        row = table_events.add_row()
        bg = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(data):
            cell = row.cells[i]
            cell.text = text
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            if i == 0:
                p.runs[0].font.bold = True

    for row in table_events.rows:
        for i, w in enumerate(ev_col_widths):
            row.cells[i].width = w

    doc.add_paragraph()

    # Section 6: Roadmap
    h6 = doc.add_heading("6. 10-Part Implementation Roadmap", level=1)
    h6.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    roadmap_steps = [
        ("Part 1: Project Setup & Docker", "Scaffold backend & frontend, configure docker-compose for Postgres & Redis, setup Express /api/health."),
        ("Part 2: Pure SQL Database Layer", "PostgreSQL schema, connection pooling with pg driver, seed scripts and parameterized query helper."),
        ("Part 3: Authentication & Security", "Bcrypt password hashing, JWT issue & middleware, Zod schemas, auth controller logic."),
        ("Part 4: Socket.io & Redis Adapter", "WebSocket server attachment, Redis adapter multi-server scaling, JWT handshake auth, useSocket hook."),
        ("Part 5: Room Management", "REST APIs and SQL queries for creating, joining, leaving rooms with real-time broadcast notifications."),
        ("Part 6: Real-Time Messaging & History", "Bidirectional message pipeline, Postgres persistence, paginated history queries."),
        ("Part 7: Presence & Typing (Redis)", "Redis Set tracking for online users, multi-tab presence, debounced typing indicator."),
        ("Part 8: React 19 Frontend UI", "AuthContext, protected routing, chat sidebar, live chat feed, auto-scrolling message list."),
        ("Part 9: Error Resilience & Testing", "Centralized Express error handling, socket reconnection hooks, load testing validation."),
        ("Part 10: Production Build & Interview Handbook", "Multi-stage Docker, environment configs, comprehensive placement interview Q&A.")
    ]

    for part, desc in roadmap_steps:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(f"{part}: ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        p.add_run(desc)

    # Section 7: Interview Q&A Cheatsheet
    h7 = doc.add_heading("7. Campus Placement Interview Preparation", level=1)
    h7.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    qa_list = [
        (
            "Q1: Why did you choose raw SQL over an ORM like Prisma or TypeORM?",
            "Answer: ORMs introduce abstraction overhead, hide underlying query behavior, and frequently produce N+1 query problems in high-throughput applications. Writing raw parameterized SQL demonstrates understanding of relational indexes, connection pooling, complex joins, and SQL injection defense directly at the database driver level."
        ),
        (
            "Q2: How does Socket.io scale horizontally across multiple servers?",
            "Answer: By default, WebSockets maintain stateful in-memory connections to a single server instance. If Server A hosts User 1 and Server B hosts User 2 in the same room, Server A cannot directly emit to User 2. The Socket.io Redis Adapter solves this by using Redis Pub/Sub channels to broadcast events across all server instances."
        ),
        (
            "Q3: How do you handle the multi-tab presence edge case in Redis?",
            "Answer: If a user opens 3 browser tabs, they generate 3 separate socket IDs. If we tracked user presence as a simple boolean, closing one tab would prematurely mark the user offline. Instead, we store socket IDs in a Redis Set (presence:user:<id>:sockets). When a socket disconnects, we remove only that socket ID. We mark the user offline only when the set size reaches zero."
        )
    ]

    for q, a in qa_list:
        doc.add_heading(q, level=2)
        p_ans = doc.add_paragraph(a)
        p_ans.style.font.size = Pt(10.5)

    output_path = "docs/RealTime_Team_Chat_Project_Documentation.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_document()
